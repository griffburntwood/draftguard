"""Synthetic end-to-end regressions; participant data stays outside Git."""
from io import BytesIO
import json
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook

from backend.main import app
from backend.classification import classify_email
from backend.documents import parse_document
from backend.batch import ParticipantInbox, run_inbox
from backend.models import EmailRecord

client = TestClient(app)
TEXT = ('Shipper: Example Exporter\nConsignee: Example Importer\nNotify Party: Example Importer\n'
        'Load Port: Singapore\nDischarge Port: Mersin\nContainer Count: 3\nGross Weight: 22,000 KG\n')


def email(body='Please compare the BL against the SI.', subject=''):
    return {'email_id': 'synthetic_001', 'from': 'ops@example.com', 'subject': subject,
            'body': body, 'attachments': ['attachments/test_SI.txt', 'attachments/test_BL.txt']}


def upload(bl=TEXT, si=TEXT, bl_name='test_BL.txt'):
    files = {'si_file': ('test_SI.txt', si.encode()), 'bl_file': (bl_name, bl.encode())}
    return client.post('/process/files', data={'email': json.dumps(email())}, files=files)


@pytest.mark.parametrize('body,subject,category', [
    ('Daily berthing report. Loading completed. No action required.', 'Review draft BL', 'GENERAL'),
    ('Please find attached shipping instructions.', '', 'SI_REQUEST'),
    ('Attached SI and draft BL for checking.', '', 'BL_COMPARISON'),
    ('Please send the draft BL for checking.', '', 'BL_COMPARISON'),
    ('Do not check the BL.', '', 'GENERAL'),
    ('Your mailbox has exceeded its storage limit. Verify your account.', '', 'SPAM'),
    ('Please explain the invoice charges.\nRegards\nPlease review BL', '', 'INVOICE_QUERY'),
    ('Please review the BL.\nOn Monday someone wrote:\nPlease send the SI.', '', 'BL_COMPARISON'),
])
def test_real_phrasing_patterns_with_synthetic_examples(body, subject, category):
    assert classify_email(EmailRecord.model_validate(email(body, subject))).category.value == category


@pytest.mark.parametrize('suffix', ['txt', 'docx', 'xlsx'])
def test_readers_preserve_values_and_source_locations(suffix):
    if suffix == 'txt':
        data = TEXT.encode()
    else:
        output = BytesIO()
        if suffix == 'docx':
            document = Document()
            table = document.add_table(rows=0, cols=2)
            for line in TEXT.splitlines():
                label, value = line.split(': ', 1)
                cells = table.add_row().cells
                cells[0].text, cells[1].text = label, value
            document.save(output)
        else:
            workbook = Workbook()
            for line in TEXT.splitlines():
                workbook.active.append(line.split(': ', 1))
            workbook.save(output)
        data = output.getvalue()
    parsed = parse_document(data, f'document.{suffix}', 'doc_1')
    assert parsed.issue is None
    assert parsed.record.container_count.normalized_value == 3
    assert parsed.record.gross_weight_kg.normalized_value == 22000
    assert all(getattr(parsed.record, field).state.value == 'extracted' for field in (
        'shipper', 'consignee', 'notify_party', 'port_of_loading', 'port_of_discharge'))
    assert suffix.upper() in parsed.record.shipper.evidence[0].location


def test_pdf_text_and_scan_without_extra_test_dependencies():
    # Minimal one-page PDF fixture with a correct byte-offset cross-reference table.
    def pdf(lines):
        stream = 'BT /F1 10 Tf 50 750 Td 14 TL ' + ' '.join(f'({line}) Tj T*' for line in lines) + ' ET'
        objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
                   b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
                   b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
                   f'<< /Length {len(stream)} >>\nstream\n{stream}\nendstream'.encode()]
        output = b'%PDF-1.4\n'; offsets = [0]
        for n, obj in enumerate(objects, 1):
            offsets.append(len(output)); output += f'{n} 0 obj\n'.encode() + obj + b'\nendobj\n'
        start = len(output)
        output += b'xref\n0 6\n0000000000 65535 f \n' + b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
        return output + f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()
    result = parse_document(pdf(TEXT.splitlines()), 'sample.pdf', 'pdf_1')
    assert result.record.gross_weight_kg.normalized_value == 22000
    assert 'page 1' in result.record.shipper.evidence[0].location
    assert parse_document(pdf([]), 'scan.pdf', 'pdf_2').issue == 'unreadable'


def test_spreadsheet_formulas_are_not_invented_values():
    workbook = Workbook(); workbook.active.append(['Shipper', '=A2'])
    output = BytesIO(); workbook.save(output)
    field = parse_document(output.getvalue(), 'sample.xlsx', 'sheet').record.shipper
    assert field.state.value == 'unreadable' and field.normalized_value is None


@pytest.mark.parametrize('raw,expected', [('22 MT', 22000), ('22,000 KGS', 22000), ('22 tonnes', 22000), ('22 lb', None), ('-1 KG', None), ('n/a', None)])
def test_weight_units_are_explicit(raw, expected):
    result = parse_document(f'Gross Weight: {raw}'.encode(), 'weight.txt', 'weight')
    assert result.record.gross_weight_kg.normalized_value == expected


def test_upload_and_export_preserve_defect_with_uncertainty():
    response = upload(TEXT.replace('Count: 3', 'Count: 4').replace('Gross Weight: 22,000 KG\n', ''))
    assert response.status_code == 200, response.text
    result = response.json()
    assert result['submission_entry'] == {'category': 'BL_COMPARISON', 'status': 'NEEDS_REVIEW',
        'review_reason': 'missing_value', 'has_defect': True, 'defect_fields': ['container_count']}
    assert len(result['document_hashes']['si']) == 64


@pytest.mark.parametrize('text,name,reason', [('COMMERCIAL INVOICE\n'+TEXT, 'invoice.txt', 'wrong_doc_type'), ('broken PDF', 'corrupt.pdf', 'unreadable')])
def test_invalid_documents_reach_review(text, name, reason):
    result = upload(text, bl_name=name)
    assert result.status_code == 200
    assert result.json()['submission_entry']['review_reason'] == reason


def test_upload_missing_and_unsupported():
    missing = client.post('/process/files', data={'email': json.dumps(email())})
    assert missing.json()['submission_entry']['review_reason'] == 'missing_attachment'
    assert upload(bl_name='source.exe').status_code == 415
    assert client.post('/process/files', data={'email': '{}'}).status_code == 422


def test_upload_size_limit():
    response = upload('x' * (10 * 1024 * 1024 + 1))
    assert response.status_code == 413


def test_human_correction_preserves_original_source_and_audit():
    original = upload(TEXT.replace('Count: 3', 'Count: 4')).json()
    request = {'original': original, 'reviewer': 'Test Reviewer', 'reason': 'Checked carrier revision',
               'corrections': [{'side': 'bl', 'field': 'container_count', 'value': 3}]}
    response = client.post('/process/review', json=request)
    assert response.status_code == 200, response.text
    corrected = response.json()
    assert corrected['comparison']['status'] == 'OK'
    count = next(f['bl'] for f in corrected['comparison']['field_results'] if f['field'] == 'container_count')
    assert count['raw_value'] == '4' and count['normalized_value'] == 3
    assert len(count['evidence']) == 2
    assert corrected['review_audit'][0]['previous_value'] == 4
    assert corrected['submission_entry']['has_defect'] is False
    assert corrected['document_hashes'] == original['document_hashes']


@pytest.mark.parametrize('field,value', [('shipper', '  '), ('container_count', 1.5), ('container_count', -1), ('gross_weight_kg', -1)])
def test_invalid_corrections_are_rejected(field, value):
    response = client.post('/process/review', json={'original': upload().json(), 'reviewer': 'Reviewer', 'reason': 'Checked source', 'corrections': [{'side': 'bl', 'field': field, 'value': value}]})
    assert response.status_code == 422


def test_wrong_document_cannot_be_fixed_by_a_value_override():
    original = upload('COMMERCIAL INVOICE\n'+TEXT).json()
    response = client.post('/process/review', json={'original': original, 'reviewer': 'Reviewer', 'reason': 'Override', 'corrections': [{'side': 'bl', 'field': 'shipper', 'value': 'Example'}]})
    assert response.status_code == 422


def test_participant_batch_uses_only_inbox_and_referenced_files(tmp_path):
    path = tmp_path / 'participant.zip'
    with ZipFile(path, 'w') as archive:
        archive.writestr('inbox/email_001.json', json.dumps(email()))
        archive.writestr('attachments/test_SI.txt', TEXT)
        archive.writestr('attachments/test_BL.txt', TEXT.replace('Count: 3', 'Count: 4'))
        archive.writestr('ground_truth.json', 'must not be opened')
        archive.writestr('loader.py', 'raise RuntimeError("must not run")')
    entries, details, summary = run_inbox(path)
    assert list(entries) == ['synthetic_001']
    assert entries['synthetic_001']['defect_fields'] == ['container_count']
    assert summary['emails'] == 1
    response = client.post('/process/inbox', files={'bundle': ('participant.zip', path.read_bytes())})
    assert response.status_code == 200
    assert response.json()['submission'] == entries
    inbox = ParticipantInbox(path)
    try:
        with pytest.raises(ValueError):
            inbox.read('../outside.txt')
    finally:
        inbox.close()
    assert client.post('/process/inbox', files={'bundle': ('bad.zip', b'bad')}).status_code == 422


def test_conflicting_port_codes_cannot_be_normalized_away():
    result = upload(TEXT.replace('Singapore', 'Singapore (SGXXX)'), si=TEXT.replace('Singapore', 'Singapore (SGSIN)'))
    assert result.json()['comparison']['defect_fields'] == ['port_of_loading']
