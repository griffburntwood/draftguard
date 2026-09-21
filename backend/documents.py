"""Bounded local document readers. Documents are data, never executable code."""
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile

from .extraction import LABELS, empty_record, extract_text_document, split_label
from .models import ShipmentRecord

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000
MAX_ARCHIVE_BYTES = 40 * 1024 * 1024
SUPPORTED_TYPES = {'.txt', '.pdf', '.docx', '.xlsx'}


@dataclass
class ParsedDocument:
    record: ShipmentRecord
    kind: str | None
    issue: str | None = None
    content_hash: str = ''


def document_kind(text):
    # Inspect document headings, not filenames or arbitrary mentions in the body.
    head = '\n'.join(text.splitlines()[:4]).upper()
    if any(title in head for title in ('COMMERCIAL INVOICE', 'CERTIFICATE OF ORIGIN', 'PACKING LIST')):
        return 'OTHER'
    if 'SHIPPING INSTRUCTION' in head or 'BILL OF LADING INSTRUCTION' in head:
        return 'SI'
    if 'BILL OF LADING' in head or 'DRAFT B/L' in head:
        return 'BL'
    return None


def _check_zip(data):
    with ZipFile(BytesIO(data)) as archive:
        if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > MAX_ARCHIVE_BYTES:
            raise ValueError('Document archive exceeds processing limits')


def read_document(data: bytes, filename: str):
    suffix = Path(filename).suffix.lower()
    lines, locations = [], []
    character_count = 0

    def add(text, location):
        nonlocal character_count
        character_count += len(str(text))
        for line in str(text).splitlines() or ['']:
            lines.append(line)
            locations.append(location)
        if len(lines) > 30000 or character_count > MAX_TEXT_CHARS:
            raise ValueError('Document is too large')

    if len(data) > MAX_FILE_BYTES:
        raise ValueError('Document exceeds 10 MiB')
    if suffix == '.txt':
        text = data.decode('utf-8-sig')
        lines = text.splitlines()
        locations = [f'TXT: line {n + 1}' for n in range(len(lines))]
    elif suffix == '.pdf':
        import pdfplumber
        with pdfplumber.open(BytesIO(data)) as pdf:
            if len(pdf.pages) > 50:
                raise ValueError('PDF exceeds 50 pages')
            for page in pdf.pages:
                text = page.extract_text() or ''
                for number, line in enumerate(text.splitlines(), 1):
                    # Convert known layout labels without colons to the text reader's format.
                    for label in sorted((x for values in LABELS.values() for x in values), key=len, reverse=True):
                        if not split_label(line)[0] and line.casefold().startswith(label.casefold() + ' '):
                            line = label + ': ' + line[len(label):].strip()
                            break
                    # Prevent shipment metadata/table rows becoming part of party addresses.
                    if re.match(r'(?i)^(ocean vessel|vessel|export carrier|container no\.|hs code)\b', line) and ':' not in line:
                        line = line.split(' ', 1)[0] + ': ' + line
                    add(line, f'PDF: page {page.page_number}, line {number}')
                add('', f'PDF: page {page.page_number}')
    elif suffix == '.docx':
        from docx import Document
        from docx.table import Table
        _check_zip(data)
        document = Document(BytesIO(data))
        paragraph_number = table_number = 0
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                table_number += 1
                for row_number, row in enumerate(block.rows, 1):
                    cells = [cell.text for cell in row.cells]
                    if len(cells) >= 2:
                        label = cells[0].strip().rstrip(':')
                        add(label + ': ' + '\n'.join(cells[1:]), f'DOCX: table {table_number}, row {row_number}')
            else:
                paragraph_number += 1
                add(block.text, f'DOCX: paragraph {paragraph_number}')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        _check_zip(data)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > 20:
                raise ValueError('Too many worksheets')
            for sheet in workbook:
                if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 100:
                    raise ValueError('Worksheet exceeds processing limits')
                for row in sheet:
                    nonempty = [cell for cell in row if cell.value is not None]
                    if not nonempty:
                        continue
                    label = str(nonempty[0].value).strip().rstrip(':')
                    values = ['[formula unavailable]' if cell.data_type == 'f' else str(cell.value) for cell in nonempty[1:]]
                    add(label + (': ' + ' '.join(values) if values else ''), f'XLSX: {sheet.title}!{nonempty[0].coordinate}:{nonempty[-1].coordinate}')
                add('', f'XLSX: {sheet.title}')
        finally:
            workbook.close()
    else:
        raise ValueError('Unsupported file type')
    text = '\n'.join(lines)
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError('Extracted text exceeds processing limits')
    if not text.strip():
        raise ValueError('No readable text; scanned documents require manual review')
    return text, locations


def parse_document(data: bytes, filename: str, document_id: str, expected_role: str | None = None) -> ParsedDocument:
    digest = sha256(data).hexdigest()
    try:
        text, locations = read_document(data, filename)
    except Exception:
        # Corrupt/encrypted/unsupported/scan-only inputs are expected review cases.
        return ParsedDocument(empty_record(document_id), None, 'unreadable', digest)
    kind = document_kind(text)
    if expected_role and kind and kind != expected_role:
        return ParsedDocument(empty_record(document_id), kind, 'wrong_doc_type', digest)
    record = extract_text_document(text, document_id, locations=None if Path(filename).suffix.lower() == '.txt' else locations)
    return ParsedDocument(record, kind, None, digest)


def extract_document(path, document_id) -> ShipmentRecord:
    path = Path(path)
    try:
        with path.open('rb') as stream:
            data = stream.read(MAX_FILE_BYTES + 1)
        return parse_document(data, path.name, document_id).record
    except OSError:
        return empty_record(document_id)
